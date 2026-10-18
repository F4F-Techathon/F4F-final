from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def rows(sheet):
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def evidence_id(shop_uuid: str, sheet: str, row: int, content: str) -> str:
    digest = hashlib.sha1(f"{shop_uuid}|{sheet}|{row}|{content}".encode("utf-8")).hexdigest()[:12]
    return f"ev_{digest}"


def sentence_quote(content: str, keywords: list[str], limit: int = 68) -> str:
    cleaned = re.sub(r"\[[^\]]+\]", "", content).replace("\n", " ").strip()
    pieces = [piece.strip() for piece in re.split(r"[。！？!?]", cleaned) if piece.strip()]
    selected = next((piece for piece in pieces if any(word in piece for word in keywords)), pieces[0] if pieces else cleaned)
    return selected[:limit]


def make_evidence(restaurant_id, shop_uuid, sheet, row, content, sentiment, keywords, tags, is_demo=False):
    quote = sentence_quote(content, keywords)
    start = content.find(quote)
    if start < 0:
        compact = content.replace("\n", " ")
        start = compact.find(quote)
        content = compact
    start = max(0, start)
    return {
        "id": evidence_id(shop_uuid, sheet, row, content),
        "restaurantId": restaurant_id,
        "sourceSheet": sheet,
        "sourceRow": row,
        "sourceText": content,
        "quote": quote,
        "highlightStart": start,
        "highlightEnd": start + len(quote),
        "sentiment": sentiment,
        "tags": tags,
        "isDemo": is_demo,
    }


def fake_restaurant(item):
    evidences = []
    for index, source in enumerate(item.pop("sources"), start=2):
        evidences.append(
            make_evidence(
                item["id"],
                item["shopUuid"],
                "模拟评论",
                index,
                source["text"],
                source["sentiment"],
                source["keywords"],
                source["tags"],
                True,
            )
        )
    item["evidence"] = evidences
    return item


def build(source: Path):
    workbook = load_workbook(source, read_only=True, data_only=True)
    shop = rows(workbook["商户基础信息"])[0]
    dishes = rows(workbook["推荐菜"])
    latest = rows(workbook["最新评论"])
    bad = rows(workbook["差评"])
    restaurant_id = "dxx-ningbo"
    shop_uuid = str(shop["shop_uuid"])

    dish_names = []
    for dish in dishes:
        name = dish.get("dish_name") or dish.get("name") or dish.get("recommend_name")
        if name and name not in dish_names:
            dish_names.append(str(name))
    if not dish_names:
        dish_names = ["梭子蟹炒年糕", "葱香元宝虾", "海肠捞饭"]

    evidences = []
    positive_keywords = ["服务", "环境", "新鲜", "梭子蟹", "元宝虾"]
    picked = []
    for row_index, review in enumerate(latest, start=2):
        content = str(review.get("content") or "").strip()
        star = float(review.get("star") or 0)
        if star >= 4.5 and len(content) >= 30 and any(word in content for word in positive_keywords):
            picked.append((row_index, content))
        if len(picked) == 4:
            break
    for row_index, content in picked:
        evidences.append(make_evidence(restaurant_id, shop_uuid, "最新评论", row_index, content, "positive", positive_keywords, ["口味", "环境", "服务"]))
    for row_index, review in enumerate(bad, start=2):
        content = str(review.get("content") or "").strip()
        evidences.append(
            make_evidence(
                restaurant_id,
                shop_uuid,
                "差评",
                row_index,
                content,
                "negative",
                ["上菜", "海肠", "服务", "刺鼻", "雷达", "一般"],
                ["服务", "出餐", "菜品"],
            )
        )

    score_text = str(shop.get("score") or "")
    rating_match = re.search(r"口味\s*([0-9.]+)", score_text)
    real = {
        "id": restaurant_id,
        "shopUuid": shop_uuid,
        "agentId": "hua",
        "assetId": "hua",
        "name": str(shop["shop_name"]),
        "shortName": "丁香小鲜",
        "cuisine": "宁波菜／浙菜／海鲜",
        "district": f"{shop.get('district', '')}·{shop.get('region', '')}",
        "address": str(shop.get("address") or ""),
        "avgPrice": int(float(shop.get("avg_price") or 0)),
        "rating": float(rating_match.group(1)) if rating_match else 4.6,
        "hours": str(shop.get("business_time") or ""),
        "recommendedDishes": dish_names[:8],
        "highlights": ["活鲜现做", "包间选择多", "适合聚餐和商务宴请"],
        "cautions": ["以海鲜菜为主", "评论对口味浓淡评价不一", "少数评论提到服务和出餐问题"],
        "isDemo": False,
        "evidence": evidences,
    }

    rong = fake_restaurant(
        {
            "id": "rong-hotpot",
            "shopUuid": "DEMO-RONG-001",
            "agentId": "rong",
            "assetId": "rong",
            "name": "蓉姐·巷子川味",
            "shortName": "蓉姐川味",
            "cuisine": "川菜／小火锅",
            "district": "浦东新区·张江",
            "address": "演示数据：张江路88号",
            "avgPrice": 86,
            "rating": 4.5,
            "hours": "演示数据：11:00-22:30",
            "recommendedDishes": ["藤椒牛肉", "双椒鱼片", "不辣番茄小锅"],
            "highlights": ["价格友好", "可调辣度", "适合朋友聚餐"],
            "cautions": ["部分蘸料含花生和坚果", "高峰期可能排队"],
            "isDemo": True,
            "sources": [
                {"text": "演示评论：藤椒牛肉香而不燥，辣度可以调整，三个人点小份也够吃。", "sentiment": "positive", "keywords": ["藤椒牛肉", "辣度"], "tags": ["口味", "分量"]},
                {"text": "演示评论：周五晚上等位二十多分钟，蘸料台有花生碎，过敏的人要提前说明。", "sentiment": "negative", "keywords": ["等位", "花生"], "tags": ["排队", "过敏"]},
            ],
        }
    )
    pierre = fake_restaurant(
        {
            "id": "pierre-bistro",
            "shopUuid": "DEMO-PIERRE-001",
            "agentId": "pierre",
            "assetId": "pierre",
            "name": "Pierre 的晚餐小馆",
            "shortName": "Pierre小馆",
            "cuisine": "法餐／西餐",
            "district": "浦东新区·张江",
            "address": "演示数据：科苑路18号",
            "avgPrice": 198,
            "rating": 4.7,
            "hours": "演示数据：17:00-23:00",
            "recommendedDishes": ["香草烤鸡", "红酒炖牛肉", "无酒精莓果气泡饮"],
            "highlights": ["环境安静", "适合约会", "可做无酒精搭配"],
            "cautions": ["部分酱汁含乳制品", "预算相对较高"],
            "isDemo": True,
            "sources": [
                {"text": "演示评论：灯光安静，座位距离舒服，香草烤鸡火候很好，适合两个人慢慢吃。", "sentiment": "positive", "keywords": ["灯光", "香草烤鸡"], "tags": ["环境", "口味"]},
                {"text": "演示评论：套餐价格偏高，奶油酱汁比较多，对乳制品过敏需要提前沟通。", "sentiment": "negative", "keywords": ["价格", "乳制品"], "tags": ["价格", "过敏"]},
            ],
        }
    )
    return {
        "schemaVersion": 1,
        "source": {"type": "xlsx", "file": source.name, "note": "丁香小鲜来自源工作簿；另外两家为明确标注的演示数据。"},
        "restaurants": [real, rong, pierre],
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: python scripts/import_restaurants.py source.xlsx data/restaurants.json")
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(build(Path(sys.argv[1])), ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
